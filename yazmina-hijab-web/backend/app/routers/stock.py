"""
Yazmina Hijab Web — Stock Manager API.

Provides SKU listing and Excel export for BigSeller import (stock IN/OUT).
"""
import io
from datetime import datetime
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..database import get_db
from ..auth import get_current_admin
from ..models.admin_user import AdminUser
from ..models.sku import SkuMaster

router = APIRouter(prefix="/api/stock", tags=["stock"])


# ── SKU List ──────────────────────────────────────────────────────────

@router.get("/skus")
def list_skus(
    search: str = "",
    db: Session = Depends(get_db),
    admin: AdminUser = Depends(get_current_admin),
):
    """List active SKUs for the stock manager dropdown."""
    q = db.query(SkuMaster).filter(SkuMaster.is_active == 1)
    if search:
        q = q.filter(
            SkuMaster.kode_sku.ilike(f"%{search}%")
            | SkuMaster.nama_produk.ilike(f"%{search}%")
        )
    skus = q.order_by(SkuMaster.kode_sku).limit(200).all()
    return [
        {
            "id": s.id,
            "kode_sku": s.kode_sku,
            "nama_produk": s.nama_produk,
            "harga_jual": s.harga_jual,
            "kategori": s.kategori,
        }
        for s in skus
    ]


# ── Export Excel ──────────────────────────────────────────────────────

class StockItem(BaseModel):
    sku: str
    qty: int
    harga: Optional[float] = 0


class ExportRequest(BaseModel):
    items: List[StockItem]
    mode: str  # "stock_in" or "stock_out"


@router.post("/export")
def export_stock(
    body: ExportRequest,
    admin: AdminUser = Depends(get_current_admin),
):
    """
    Generate BigSeller-compatible Excel file for stock IN or OUT.
    Returns the file as a download.
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl tidak terinstall")

    if not body.items:
        raise HTTPException(status_code=400, detail="Daftar item kosong")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stock Import"

    # Styling
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="00F0FF", end_color="00F0FF", fill_type="solid")
    header_font_dark = Font(bold=True, size=11, color="000000")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    if body.mode == "stock_in":
        # BigSeller Stock IN format
        headers = [
            "*Nomor SKU\n(SKU atau GTIN Wajib Diisi)",
            "*GTIN\n(SKU atau GTIN Wajib Diisi)",
            "*Jumlah Penambahan Stok",
            "Harga Satuan",
            "Tanggal Produksi",
            "Tanggal Kedaluwarsa",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_dark
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row_idx, item in enumerate(body.items, 2):
            ws.cell(row=row_idx, column=1, value=item.sku).border = thin_border
            ws.cell(row=row_idx, column=2, value="").border = thin_border
            ws.cell(row=row_idx, column=3, value=item.qty).border = thin_border
            ws.cell(row=row_idx, column=4, value=item.harga if item.harga > 0 else "").border = thin_border
            ws.cell(row=row_idx, column=5, value="").border = thin_border
            ws.cell(row=row_idx, column=6, value="").border = thin_border

        # Column widths
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 20
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 15

        filename = f"Penambahan_Stok_{datetime.today().strftime('%Y%m%d')}.xlsx"

    else:
        # BigSeller Stock OUT format
        headers = [
            "*Nomor SKU",
            "*Jumlah Pengurangan Stok",
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font_dark
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border

        for row_idx, item in enumerate(body.items, 2):
            ws.cell(row=row_idx, column=1, value=item.sku).border = thin_border
            ws.cell(row=row_idx, column=2, value=item.qty).border = thin_border

        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 25

        filename = f"Pengurangan_Stok_{datetime.today().strftime('%Y%m%d')}.xlsx"

    # Save to buffer
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/export-preview")
def export_preview(
    items: str = "",
    mode: str = "stock_in",
    admin: AdminUser = Depends(get_current_admin),
):
    """
    Preview what the Excel will look like (JSON format).
    items is a JSON array string: [{"sku":"DG-L","qty":10,"harga":85000}]
    """
    import json
    try:
        item_list = json.loads(items) if items else []
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="Invalid items JSON")

    if mode == "stock_in":
        columns = ["SKU", "GTIN", "Qty", "Harga", "ProdDate", "ExpDate"]
        rows = [
            [it.get("sku", ""), "", it.get("qty", 0), it.get("harga", 0), "", ""]
            for it in item_list
        ]
    else:
        columns = ["SKU", "Qty"]
        rows = [[it.get("sku", ""), it.get("qty", 0)] for it in item_list]

    return {"columns": columns, "rows": rows, "mode": mode}
