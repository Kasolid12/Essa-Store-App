"""
Yazmina Hijab Web — Gaji (Salary) CRUD API.

Features:
- CRUD salary runs with bon integration
- Import Excel penjahit (garapan lines from Excel)
- Import Excel pengsup (setoran from Excel)
- Import Excel absensi karyawan (fingerprint machine Excel)
- Batch save pasukan karyawan
"""

import io
import os
import re
from typing import Optional, List

import pandas as pd
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..database import get_db
from ..auth import get_current_admin
from ..models.payroll import (
    SalaryRun, SalaryLineItem, BonBalance, BonMovement,
    AttendanceRecord, MasterTarifPenjahit,
)
from ..models.person import Person
from ..models.sku import SkuMaster

router = APIRouter(prefix="/api/gaji", tags=["gaji"])

TIPE_OPTIONS = ["BORONGAN_PENJAHIT", "PENGSUP", "PASUKAN_KARYAWAN"]


# ── Schemas ───────────────────────────────────────────────────────────

class LineItemCreate(BaseModel):
    sku_id: Optional[int] = None
    model_code: Optional[str] = None
    qty: int = 1
    tarif_per_pcs: float = 0.0


class SalaryRunCreate(BaseModel):
    tipe: str
    person_id: Optional[int] = None
    periode_mulai: Optional[str] = None
    periode_akhir: Optional[str] = None
    tanggal_proses: str
    tambah_bon: float = 0.0
    potong_bon: float = 0.0
    catatan: Optional[str] = None
    line_items: List[LineItemCreate] = []


class SalaryRunUpdate(BaseModel):
    tipe: Optional[str] = None
    person_id: Optional[int] = None
    periode_mulai: Optional[str] = None
    periode_akhir: Optional[str] = None
    tanggal_proses: Optional[str] = None
    tambah_bon: Optional[float] = None
    potong_bon: Optional[float] = None
    catatan: Optional[str] = None


class SalaryRunOut(BaseModel):
    id: int
    tipe: str
    person_id: Optional[int] = None
    person_nama: Optional[str] = None
    periode_mulai: Optional[str] = None
    periode_akhir: Optional[str] = None
    tanggal_proses: str
    gaji_kotor: float
    bon_lama: float
    tambah_bon: float
    potong_bon: float
    gaji_bersih: float
    sisa_bon_akhir: float
    catatan: Optional[str] = None
    is_deleted: int
    line_items_count: int = 0

    model_config = {"from_attributes": True}


# ── Karyawan Attendance Schema ────────────────────────────────────────

class KaryawanRow(BaseModel):
    person_id: int
    nama: str
    hadir: int = 0
    menit_normal: float = 0
    tarif_normal: float = 140.0
    menit_lembur: float = 0
    tarif_lembur: float = 160.0
    gaji_kotor: float = 0
    bon_lama: float = 0
    potong_bon: float = 0
    gaji_bersih: float = 0
    daily_records: List[dict] = []


class SavePasukanRequest(BaseModel):
    tanggal_proses: str
    tarif_normal: float = 140.0
    tarif_lembur: float = 160.0
    karyawan: List[KaryawanRow]


class DailyEdit(BaseModel):
    tanggal: str
    tap_masuk: str = ""
    tap_keluar: str = ""
    menit_normal: float = 0
    menit_lembur: float = 0


class EditKaryawanRequest(BaseModel):
    person_id: int
    tarif_normal: float = 140.0
    tarif_lembur: float = 160.0
    potong_bon: float = 0
    daily_records: List[DailyEdit]


# ── Helpers ───────────────────────────────────────────────────────────

def _get_or_create_balance(person_id: int, db: Session) -> BonBalance:
    balance = db.query(BonBalance).filter(BonBalance.person_id == person_id).first()
    if not balance:
        balance = BonBalance(person_id=person_id, saldo=0.0)
        db.add(balance)
        db.flush()
    return balance


def _enrich(run: SalaryRun, db: Session) -> dict:
    person = db.get(Person, run.person_id) if run.person_id else None
    return {
        "id": run.id,
        "tipe": run.tipe,
        "person_id": run.person_id,
        "person_nama": person.nama if person else None,
        "periode_mulai": run.periode_mulai,
        "periode_akhir": run.periode_akhir,
        "tanggal_proses": run.tanggal_proses,
        "gaji_kotor": run.gaji_kotor,
        "bon_lama": run.bon_lama,
        "tambah_bon": run.tambah_bon,
        "potong_bon": run.potong_bon,
        "gaji_bersih": run.gaji_bersih,
        "sisa_bon_akhir": run.sisa_bon_akhir,
        "catatan": run.catatan,
        "is_deleted": run.is_deleted,
        "line_items_count": len(run.line_items),
    }


def _recalc(run: SalaryRun):
    """Recalculate salary totals from line items."""
    run.gaji_kotor = sum(item.subtotal for item in run.line_items)
    run.gaji_bersih = run.gaji_kotor - run.potong_bon
    run.sisa_bon_akhir = run.bon_lama + run.tambah_bon - run.potong_bon


def _sync_bon_on_create(run: SalaryRun, db: Session):
    """Sync BonBalance + create BonMovement when creating salary run."""
    if not run.person_id:
        return

    balance = _get_or_create_balance(run.person_id, db)
    tgl = run.tanggal_proses
    source = f"PAYROLL_{run.tipe}"

    if run.tambah_bon > 0:
        balance.saldo += run.tambah_bon
        db.add(BonMovement(
            person_id=run.person_id, tanggal=tgl, tipe="TAMBAH",
            nominal=run.tambah_bon, sumber=source,
        ))

    if run.potong_bon > 0:
        balance.saldo -= run.potong_bon
        db.add(BonMovement(
            person_id=run.person_id, tanggal=tgl, tipe="POTONG_GAJI",
            nominal=run.potong_bon, sumber=source,
        ))


def _reverse_bon(old_run: SalaryRun, db: Session):
    """Reverse previous bon changes (for update/delete)."""
    if not old_run.person_id:
        return

    balance = db.query(BonBalance).filter(BonBalance.person_id == old_run.person_id).first()
    if not balance:
        return

    if old_run.tambah_bon > 0:
        balance.saldo -= old_run.tambah_bon

    if old_run.potong_bon > 0:
        balance.saldo += old_run.potong_bon

    source = f"PAYROLL_{old_run.tipe}"
    db.query(BonMovement).filter(
        BonMovement.person_id == old_run.person_id,
        BonMovement.tanggal == old_run.tanggal_proses,
        BonMovement.sumber == source,
    ).delete(synchronize_session=False)


def _parse_waktu(v) -> Optional[float]:
    """Parse time from Excel fingerprint (HH:MM or decimal 0.0-1.0)."""
    v_str = str(v).strip()
    if v_str in ['', 'nan', 'None']:
        return None
    if ":" in v_str:
        try:
            parts = v_str.split(":")
            return (int(parts[0]) * 60 + int(parts[1])) / 1440.0
        except Exception:
            return None
    try:
        val = float(v_str)
        if 0.0 <= val <= 1.0:
            return val
    except Exception:
        pass
    return None


# ── Routes ────────────────────────────────────────────────────────────

@router.get("/tipe")
def list_tipe():
    return TIPE_OPTIONS


@router.get("/bon-lama/{person_id}")
def get_bon_lama(
    person_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Get current bon balance for a person (auto-load for salary form)."""
    balance = db.query(BonBalance).filter(BonBalance.person_id == person_id).first()
    return {"person_id": person_id, "bon_lama": balance.saldo if balance else 0.0}


@router.get("")
def list_gaji(
    tipe: Optional[str] = Query(None),
    date_from: Optional[str] = Query(None),
    date_to: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    q = db.query(SalaryRun).filter(SalaryRun.is_deleted == 0)

    if tipe:
        q = q.filter(SalaryRun.tipe == tipe)
    if date_from:
        q = q.filter(SalaryRun.tanggal_proses >= date_from)
    if date_to:
        q = q.filter(SalaryRun.tanggal_proses <= date_to)

    runs = q.order_by(SalaryRun.tanggal_proses.desc(), SalaryRun.id.desc()).limit(200).all()
    return [_enrich(r, db) for r in runs]


@router.get("/{run_id}")
def get_gaji(
    run_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    run = db.get(SalaryRun, run_id)
    if not run or run.is_deleted:
        raise HTTPException(404, "Data gaji tidak ditemukan")

    result = _enrich(run, db)
    result["line_items"] = [
        {
            "id": li.id,
            "sku_id": li.sku_id,
            "model_code": li.model_code,
            "qty": li.qty,
            "tarif_per_pcs": li.tarif_per_pcs,
            "subtotal": li.subtotal,
        }
        for li in run.line_items
    ]
    return result


@router.post("")
def create_gaji(
    body: SalaryRunCreate,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    if body.tipe not in TIPE_OPTIONS:
        raise HTTPException(400, f"Tipe '{body.tipe}' tidak valid")

    bon_lama = 0.0
    if body.person_id:
        balance = db.query(BonBalance).filter(BonBalance.person_id == body.person_id).first()
        bon_lama = balance.saldo if balance else 0.0

    run = SalaryRun(
        tipe=body.tipe,
        person_id=body.person_id,
        periode_mulai=body.periode_mulai,
        periode_akhir=body.periode_akhir,
        tanggal_proses=body.tanggal_proses,
        bon_lama=bon_lama,
        tambah_bon=body.tambah_bon,
        potong_bon=body.potong_bon,
        catatan=body.catatan,
    )

    for li_data in body.line_items:
        subtotal = li_data.qty * li_data.tarif_per_pcs
        li = SalaryLineItem(
            sku_id=li_data.sku_id,
            model_code=li_data.model_code,
            qty=li_data.qty,
            tarif_per_pcs=li_data.tarif_per_pcs,
            subtotal=subtotal,
        )
        run.line_items.append(li)

    _recalc(run)
    db.add(run)
    db.flush()
    _sync_bon_on_create(run, db)
    db.commit()
    db.refresh(run)
    return _enrich(run, db)


@router.put("/{run_id}")
def update_gaji(
    run_id: int,
    body: SalaryRunUpdate,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    run = db.get(SalaryRun, run_id)
    if not run or run.is_deleted:
        raise HTTPException(404, "Data gaji tidak ditemukan")

    data = body.model_dump(exclude_unset=True)
    if "tipe" in data and data["tipe"] not in TIPE_OPTIONS:
        raise HTTPException(400, f"Tipe '{data['tipe']}' tidak valid")

    _reverse_bon(run, db)

    for k, v in data.items():
        setattr(run, k, v)

    if run.person_id:
        balance = db.query(BonBalance).filter(BonBalance.person_id == run.person_id).first()
        run.bon_lama = balance.saldo if balance else 0.0

    _recalc(run)
    db.flush()
    _sync_bon_on_create(run, db)
    db.commit()
    db.refresh(run)
    return _enrich(run, db)


@router.delete("/{run_id}")
def delete_gaji(
    run_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    run = db.get(SalaryRun, run_id)
    if not run or run.is_deleted:
        raise HTTPException(404, "Data gaji tidak ditemukan")

    _reverse_bon(run, db)
    run.is_deleted = 1
    db.commit()
    return {"message": "Data gaji dihapus"}


# ══════════════════════════════════════════════════════════════════════
# IMPORT EXCEL — PENJAHIT
# ══════════════════════════════════════════════════════════════════════

@router.post("/import-excel-penjahit")
async def import_excel_penjahit(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Import garapan penjahit from Excel. Expects columns: SKU/Garapan, Qty, Harga/Tarif (optional)."""
    content = await file.read()
    try:
        if file.filename.endswith('.csv'):
            df = pd.read_csv(io.BytesIO(content))
        else:
            df = pd.read_excel(io.BytesIO(content))
    except Exception as e:
        raise HTTPException(400, f"Gagal membaca file: {e}")

    # Normalize column names
    df.columns = df.columns.astype(str).str.strip().str.lower()

    col_sku = next((c for c in df.columns if 'sku' in c or 'garapan' in c or 'jenis' in c), None)
    col_qty = next((c for c in df.columns if 'qty' in c or 'jumlah' in c), None)
    col_harga = next((c for c in df.columns if 'harga' in c or 'tarif' in c), None)

    if not col_sku or not col_qty:
        raise HTTPException(400, "File harus memiliki kolom 'SKU/Garapan' dan 'Qty'")

    items = []
    for _, row in df.iterrows():
        nama = str(row[col_sku]).strip()
        qty = float(row[col_qty])
        harga = float(row[col_harga]) if col_harga and pd.notnull(row[col_harga]) else 0

        tarif_id = None
        if harga == 0:
            tarif_db = db.query(MasterTarifPenjahit).filter(
                MasterTarifPenjahit.kode_garapan == nama
            ).first()
            if tarif_db:
                harga = tarif_db.harga
                tarif_id = tarif_db.id

        if harga > 0 and qty > 0:
            items.append({
                "tarif_id": tarif_id,
                "sku_id": None,
                "nama_garapan": nama,
                "qty": qty,
                "harga": harga,
                "total": qty * harga,
            })

    return {"items": items, "count": len(items)}


# ══════════════════════════════════════════════════════════════════════
# IMPORT EXCEL — PENGSUP
# ══════════════════════════════════════════════════════════════════════

@router.post("/import-excel-pengsup")
async def import_excel_pengsup(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Smart Import Draft Pengsup from Excel. Spatial scanning for metadata."""
    content = await file.read()
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Gagal membaca file: {e}")

    # Find the right sheet
    if "Daftar_Pemasukan" in wb.sheetnames:
        ws = wb["Daftar_Pemasukan"]
    else:
        ws = wb.active

    kain_qty = 0.0
    kain_harga = 0.0
    tambah_bon = 0.0
    potong_bon = 0.0
    data_start_row = 8

    # Spatial scanning for metadata
    for r in range(1, min(20, ws.max_row + 1)):
        for c in range(1, min(5, ws.max_column + 1)):
            cell_val = str(ws.cell(r, c).value or "").strip().lower()
            if not cell_val:
                continue

            val_cell = ws.cell(r, c + 1)

            if "kain_qty" in cell_val or "kain qty" in cell_val:
                val_raw = str(val_cell.value or "0").replace(',', '')
                kain_qty = float(val_raw) if val_raw else 0.0
            elif "kain_harga" in cell_val or "harga kain" in cell_val:
                val_raw = str(val_cell.value or "0").replace('Rp', '').replace(',', '').strip()
                kain_harga = float(val_raw) if val_raw else 0.0
            elif "bon_tambah" in cell_val or "tambah bon" in cell_val:
                val_raw = str(val_cell.value or "0").replace('Rp', '').replace(',', '').strip()
                tambah_bon = float(val_raw) if val_raw else 0.0
            elif "bon_potong" in cell_val or "potong bon" in cell_val:
                val_raw = str(val_cell.value or "0").replace('Rp', '').replace(',', '').strip()
                potong_bon = float(val_raw) if val_raw else 0.0
            elif "sku" in cell_val or "nama barang" in cell_val or "jenis garapan" in cell_val:
                data_start_row = r

    # Read items below header
    items = []
    for r in range(data_start_row + 1, ws.max_row + 1):
        nama = ws.cell(r, 1).value
        if not nama or str(nama).strip() == "" or "total" in str(nama).lower():
            continue

        qty_val = str(ws.cell(r, 2).value or "0").replace(',', '').strip()
        harga_val = str(ws.cell(r, 3).value or "0").replace('Rp ', '').replace('.', '').replace(',', '').strip()
        tipe = str(ws.cell(r, 4).value or "Setor Barang Jadi (Kain)").strip()

        qty = float(qty_val) if qty_val else 0.0
        harga = float(harga_val) if harga_val else 0.0

        items.append({
            "tipe": tipe,
            "sku_kode": str(nama),
            "nama_garapan": str(nama),
            "qty": qty,
            "harga": harga,
            "total": qty * harga,
        })

    return {
        "items": items,
        "count": len(items),
        "kain_qty": kain_qty,
        "kain_harga": kain_harga,
        "tambah_bon": tambah_bon,
        "potong_bon": potong_bon,
    }


# ══════════════════════════════════════════════════════════════════════
# IMPORT EXCEL — ABSENSI KARYAWAN (FINGERPRINT)
# ══════════════════════════════════════════════════════════════════════

@router.post("/import-absensi")
async def import_absensi(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Import Excel from fingerprint machine. Spatial scanning for User ID, name, dates."""
    content = await file.read()
    try:
        excel_data = pd.read_excel(io.BytesIO(content), sheet_name=None, header=None)
    except Exception as e:
        raise HTTPException(400, f"Gagal membaca file absensi: {e}")

    rekap_dict = {}  # name -> {hadir, normal, lembur, daily_records}

    for sheet_name, df in excel_data.items():
        sheet_str = str(sheet_name)
        if "Analisa Kehadiran" in sheet_str or "Pengaturan Shift" in sheet_str:
            continue

        # Find "user id" anchor
        uid_loc = []
        for r in range(len(df)):
            for c in range(len(df.columns)):
                if "user id" in str(df.iloc[r, c]).strip().lower():
                    uid_loc.append((r, c))

        if not uid_loc:
            continue

        for r_uid, c_uid in uid_loc:
            emp_id = ""
            name_str = "Unknown"

            # Find employee ID
            for offset in range(1, 10):
                if c_uid + offset < len(df.columns):
                    val = str(df.iloc[r_uid, c_uid + offset]).strip()
                    if val not in ['', 'nan', 'None']:
                        emp_id = val.replace(".0", "")
                        break

            # Find employee name
            for r_nama in range(max(0, r_uid - 2), r_uid + 2):
                for c_nama in range(max(0, c_uid - 5), c_uid + 5):
                    if c_nama < len(df.columns):
                        val = str(df.iloc[r_nama, c_nama]).strip().lower()
                        if "nama" in val:
                            for off in range(1, 10):
                                if c_nama + off < len(df.columns):
                                    nv = str(df.iloc[r_nama, c_nama + off]).strip()
                                    if nv not in ['', 'nan', 'None']:
                                        name_str = nv.lower()
                                        break
                            break

            if name_str not in rekap_dict:
                rekap_dict[name_str] = {'emp_id': emp_id, 'hadir': 0, 'normal': 0, 'lembur': 0, 'daily_records': []}

            # Find date column
            date_col = -1
            for c_test in range(c_uid, -1, -1):
                for r_test in range(r_uid + 5, min(r_uid + 20, len(df))):
                    val = str(df.iloc[r_test, c_test]).strip()
                    if len(val) >= 4 and val[:2].isdigit() and " " in val:
                        date_col = c_test
                        break
                if date_col != -1:
                    break

            if date_col == -1:
                continue

            total_normal = 0
            total_lembur = 0
            hari_hadir = 0

            for r_data in range(r_uid + 5, min(r_uid + 45, len(df))):
                date_val = str(df.iloc[r_data, date_col]).strip()
                if not (len(date_val) >= 4 and date_val[:2].isdigit()):
                    continue

                waktu_taps = []
                for c in range(date_col + 1, min(date_col + 15, len(df.columns))):
                    parsed = _parse_waktu(df.iloc[r_data, c])
                    if parsed is not None:
                        waktu_taps.append(parsed)

                if waktu_taps:
                    min_t = round(min(waktu_taps) * 1440)
                    max_t = round(max(waktu_taps) * 1440)

                    jam_masuk = f"{int(min_t // 60):02d}:{int(min_t % 60):02d}"

                    if len(waktu_taps) > 1:
                        jam_keluar = f"{int(max_t // 60):02d}:{int(max_t % 60):02d}"
                        diff = max_t - min_t
                        if diff < 0:
                            diff += 1440
                        total_mnt = diff
                        hari_hadir += 1
                    else:
                        jam_keluar = "Lupa"
                        total_mnt = 0
                        hari_hadir += 1

                    menit_normal = min(total_mnt, 480)
                    lembur = max(0, total_mnt - 480)

                    total_normal += menit_normal
                    total_lembur += lembur

                    rekap_dict[name_str]['daily_records'].append({
                        "tanggal": date_val[:6],
                        "masuk": jam_masuk,
                        "keluar": jam_keluar,
                        "menit_normal": menit_normal,
                        "menit_lembur": lembur,
                    })

            rekap_dict[name_str]['hadir'] += hari_hadir
            rekap_dict[name_str]['normal'] += total_normal
            rekap_dict[name_str]['lembur'] += total_lembur

    # Match to persons in database
    karyawans = db.query(Person).filter(Person.person_type == 'KARYAWAN').order_by(Person.nama).all()
    all_balances = db.query(BonBalance).all()
    dict_balances = {b.person_id: b.saldo for b in all_balances}

    result = []
    for k in karyawans:
        nama_lower = k.nama.strip().lower()
        matched_name = None
        for ex_name in rekap_dict.keys():
            if ex_name in nama_lower or nama_lower in ex_name:
                matched_name = ex_name
                break

        if matched_name:
            data = rekap_dict[matched_name]
            bon_lama = dict_balances.get(k.id, 0.0)
            gaji_kotor = (data['normal'] * 140.0) + (data['lembur'] * 160.0)
            potong_bon = min(bon_lama, gaji_kotor)
            gaji_bersih = gaji_kotor - potong_bon

            result.append({
                "person_id": k.id,
                "nama": k.nama,
                "emp_id": data['emp_id'],
                "hadir": data['hadir'],
                "menit_normal": data['normal'],
                "tarif_normal": 140.0,
                "menit_lembur": data['lembur'],
                "tarif_lembur": 160.0,
                "gaji_kotor": gaji_kotor,
                "bon_lama": bon_lama,
                "potong_bon": potong_bon,
                "gaji_bersih": gaji_bersih,
                "daily_records": data['daily_records'],
            })
        else:
            result.append({
                "person_id": k.id,
                "nama": k.nama,
                "emp_id": "",
                "hadir": 0,
                "menit_normal": 0,
                "tarif_normal": 140.0,
                "menit_lembur": 0,
                "tarif_lembur": 160.0,
                "gaji_kotor": 0,
                "bon_lama": dict_balances.get(k.id, 0.0),
                "potong_bon": 0,
                "gaji_bersih": 0,
                "daily_records": [],
            })

    return {"karyawan": result, "total_matched": len(rekap_dict)}


# ══════════════════════════════════════════════════════════════════════
# SAVE PASUKAN (BATCH) — Simpan semua karyawan sekaligus
# ══════════════════════════════════════════════════════════════════════

@router.post("/save-pasukan")
def save_pasukan(
    body: SavePasukanRequest,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Batch save all karyawan salary runs at once (after import absensi + editing)."""
    tanggal = body.tanggal_proses
    saved = 0

    for k in body.karyawan:
        if k.menit_normal == 0 and k.menit_lembur == 0:
            continue

        sisa_bon_akhir = k.bon_lama - k.potong_bon
        keterangan = f"Hadir: {k.hadir} | Normal: {k.menit_normal:g} | Lembur: {k.menit_lembur:g}"

        # Check for existing run (dedup)
        run = db.query(SalaryRun).filter(
            SalaryRun.tipe == "PASUKAN_KARYAWAN",
            SalaryRun.person_id == k.person_id,
            SalaryRun.tanggal_proses == tanggal,
            SalaryRun.is_deleted == 0,
        ).first()

        balance = _get_or_create_balance(k.person_id, db)

        if run:
            # UPDATE existing
            db.query(SalaryLineItem).filter(SalaryLineItem.salary_run_id == run.id).delete()
            db.query(AttendanceRecord).filter(AttendanceRecord.salary_run_id == run.id).delete()
            old_potong = run.potong_bon or 0
            balance.saldo += old_potong
            db.query(BonMovement).filter(
                BonMovement.sumber == "PAYROLL_KARYAWAN",
                BonMovement.person_id == k.person_id,
                BonMovement.tanggal == tanggal,
            ).delete(synchronize_session=False)
            run.gaji_kotor = k.gaji_kotor
            run.bon_lama = k.bon_lama
            run.tambah_bon = 0
            run.potong_bon = k.potong_bon
            run.gaji_bersih = k.gaji_bersih
            run.sisa_bon_akhir = sisa_bon_akhir
            run.catatan = keterangan
        else:
            # INSERT new
            run = SalaryRun(
                tipe="PASUKAN_KARYAWAN",
                person_id=k.person_id,
                tanggal_proses=tanggal,
                gaji_kotor=k.gaji_kotor,
                bon_lama=k.bon_lama,
                tambah_bon=0,
                potong_bon=k.potong_bon,
                gaji_bersih=k.gaji_bersih,
                sisa_bon_akhir=sisa_bon_akhir,
                catatan=keterangan,
            )
            db.add(run)
        db.flush()

        # Save line items (gaji normal + lembur)
        if k.menit_normal > 0:
            db.add(SalaryLineItem(
                salary_run_id=run.id, sku_id=None,
                model_code="[GAJI_NORMAL]",
                qty=int(k.menit_normal), tarif_per_pcs=k.tarif_normal,
                subtotal=k.menit_normal * k.tarif_normal,
            ))

        if k.menit_lembur > 0:
            db.add(SalaryLineItem(
                salary_run_id=run.id, sku_id=None,
                model_code="[GAJI_LEMBUR]",
                qty=int(k.menit_lembur), tarif_per_pcs=k.tarif_lembur,
                subtotal=k.menit_lembur * k.tarif_lembur,
            ))

        # Save attendance records
        for rec in k.daily_records:
            db.add(AttendanceRecord(
                salary_run_id=run.id,
                person_id=k.person_id,
                tanggal=rec.get("tanggal", ""),
                tap_masuk=rec.get("masuk", ""),
                tap_keluar=rec.get("keluar", ""),
                menit_normal=int(rec.get("menit_normal", 0)),
                menit_lembur=int(rec.get("menit_lembur", 0)),
                tarif_normal=k.tarif_normal,
                tarif_lembur=k.tarif_lembur,
                status="NORMAL" if rec.get("keluar", "") != "Lupa" else "LUPA_TAP",
            ))

        # Update BonBalance
        if k.potong_bon > 0:
            balance.saldo -= k.potong_bon
            db.add(BonMovement(
                person_id=k.person_id, tanggal=tanggal,
                tipe="POTONG_GAJI", nominal=k.potong_bon,
                sumber="PAYROLL_KARYAWAN",
            ))

        saved += 1

    db.commit()
    return {"message": f"{saved} data gaji karyawan berhasil disimpan", "saved": saved}


# ══════════════════════════════════════════════════════════════════════
# EDIT KARYAWAN — Edit jam masuk/keluar per hari
# ══════════════════════════════════════════════════════════════════════

@router.post("/edit-karyawan")
def edit_karyawan(
    body: EditKaryawanRequest,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Recalculate karyawan salary after editing daily attendance records."""
    total_normal = 0
    total_lembur = 0
    hadir = 0

    daily_results = []
    for rec in body.daily_records:
        menit_normal = rec.menit_normal
        menit_lembur = rec.menit_lembur

        # Recalculate from tap times if both present
        if rec.tap_masuk and rec.tap_keluar and rec.tap_keluar.lower() != "lupa":
            try:
                m_h, m_m = map(int, rec.tap_masuk.split(':'))
                k_h, k_m = map(int, rec.tap_keluar.split(':'))
                min_t = m_h * 60 + m_m
                max_t = k_h * 60 + k_m
                diff = max_t - min_t
                if diff < 0:
                    diff += 1440
                menit_normal = min(diff, 480)
                menit_lembur = max(0, diff - 480)
            except Exception:
                pass

        if menit_normal > 0 or menit_lembur > 0 or (rec.tap_masuk and rec.tap_keluar.lower() != 'lupa'):
            hadir += 1

        total_normal += menit_normal
        total_lembur += menit_lembur

        daily_results.append({
            "tanggal": rec.tanggal,
            "tap_masuk": rec.tap_masuk,
            "tap_keluar": rec.tap_keluar,
            "menit_normal": menit_normal,
            "menit_lembur": menit_lembur,
        })

    gaji_kotor = (total_normal * body.tarif_normal) + (total_lembur * body.tarif_lembur)

    balance = db.query(BonBalance).filter(BonBalance.person_id == body.person_id).first()
    bon_lama = balance.saldo if balance else 0.0
    potong_bon = min(body.potong_bon, gaji_kotor)
    gaji_bersih = gaji_kotor - potong_bon
    sisa_bon_akhir = bon_lama - potong_bon

    return {
        "person_id": body.person_id,
        "hadir": hadir,
        "menit_normal": total_normal,
        "tarif_normal": body.tarif_normal,
        "menit_lembur": total_lembur,
        "tarif_lembur": body.tarif_lembur,
        "gaji_kotor": gaji_kotor,
        "bon_lama": bon_lama,
        "potong_bon": potong_bon,
        "gaji_bersih": gaji_bersih,
        "sisa_bon_akhir": sisa_bon_akhir,
        "daily_records": daily_results,
    }


# ══════════════════════════════════════════════════════════════════════
# GET ATTENDANCE RECORDS — for editor
# ══════════════════════════════════════════════════════════════════════

@router.get("/attendance/{run_id}")
def get_attendance(
    run_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """Get daily attendance records for a salary run."""
    records = db.query(AttendanceRecord).filter(
        AttendanceRecord.salary_run_id == run_id
    ).order_by(AttendanceRecord.tanggal).all()

    return [
        {
            "id": r.id,
            "person_id": r.person_id,
            "tanggal": r.tanggal,
            "tap_masuk": r.tap_masuk,
            "tap_keluar": r.tap_keluar,
            "menit_normal": r.menit_normal,
            "menit_lembur": r.menit_lembur,
            "tarif_normal": r.tarif_normal,
            "tarif_lembur": r.tarif_lembur,
            "status": r.status,
        }
        for r in records
    ]
